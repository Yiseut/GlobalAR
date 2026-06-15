#!/usr/bin/env python3
"""Apply the 2026-06-15 non-medical-aesthetic scope cleanup.

Scope definition from user feedback:
- A medical aesthetics company must have at least one medical-aesthetic
  medical device, drug, or equipment product.
- If such a company also has cosmetics or oral supplements, those can remain
  as ancillary context.
- A company with only cosmetics, skincare, oral nutrition, or oral supplements
  should not appear in final dashboard results.
"""

from __future__ import annotations

import csv
import json
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

from dashboard_scope import is_non_core_cosmetic_or_oral_product


PROJECT_DIR = Path(__file__).resolve().parents[1]
SOURCE_DIR = PROJECT_DIR.parent
DATA_DIR = PROJECT_DIR / "data"
AUDIT_DIR = DATA_DIR / "audits"
SOURCE_WORKBOOK = SOURCE_DIR / "全球医美企业库_标准化版v4.xlsx"
PRODUCT_MASTER_CSV = DATA_DIR / "product_master.csv"
TARGET_RECORD_ID = "REC_0331"
TARGET_PRODUCT_ID = "prod_facbf7fd2c9f"
ACTION = "scope_cleanup_oral_supplement_not_medical_aesthetics"
NOTE = "软删除(口服透明质酸保健品, 非医美医疗器械/药品/设备). [scope_rule_20260615]"


def norm(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def headers(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, int]:
    return {norm(cell.value): idx for idx, cell in enumerate(ws[1], start=1) if norm(cell.value)}


def ensure_cleanup_log(wb: openpyxl.Workbook) -> openpyxl.worksheet.worksheet.Worksheet:
    if "Cleanup_Log" not in wb.sheetnames:
        ws = wb.create_sheet("Cleanup_Log")
        ws.append(["timestamp", "sheet", "row_num", "row_id", "field", "old_value", "new_value", "action", "note"])
        return ws
    ws = wb["Cleanup_Log"]
    if ws.max_row < 1 or not norm(ws.cell(1, 1).value):
        ws.append(["timestamp", "sheet", "row_num", "row_id", "field", "old_value", "new_value", "action", "note"])
    return ws


def log_change(
    cleanup_ws: openpyxl.worksheet.worksheet.Worksheet,
    sheet: str,
    row_num: int | str,
    row_id: str,
    field: str,
    old_value: Any,
    new_value: Any,
    note: str = NOTE,
) -> None:
    if norm(old_value) == norm(new_value):
        return
    cleanup_ws.append(
        [
            datetime.now().isoformat(timespec="seconds"),
            sheet,
            row_num,
            row_id,
            field,
            old_value,
            new_value,
            ACTION,
            note,
        ]
    )


def set_cell(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    row: int,
    colmap: dict[str, int],
    field: str,
    value: Any,
    cleanup_ws: openpyxl.worksheet.worksheet.Worksheet,
    row_id: str,
    note: str = NOTE,
) -> bool:
    if field not in colmap:
        return False
    cell = ws.cell(row, colmap[field])
    old = cell.value
    if norm(old) == norm(value):
        return False
    cell.value = value
    log_change(cleanup_ws, ws.title, row, row_id, field, old, value, note)
    return True


def append_note(value: Any, note: str = NOTE) -> str:
    text = norm(value)
    if note in text:
        return text
    return f"{text} | {note}" if text else note


def mark_classification_layer_deleted(value: Any) -> str:
    text = norm(value)
    if not text:
        return text
    try:
        payload = json.loads(text)
    except json.JSONDecodeError:
        return text
    payload["inclusion_status"] = "deleted"
    detail = payload.get("material_taxonomy_detail")
    if isinstance(detail, dict):
        detail["inclusion_status"] = "deleted"
        detail["backfill_audit"] = append_note(detail.get("backfill_audit"))
    return json.dumps(payload, ensure_ascii=False, separators=(",", ":"))


def rows_by_field(ws: openpyxl.worksheet.worksheet.Worksheet, field: str, value: str) -> list[int]:
    colmap = headers(ws)
    if field not in colmap:
        return []
    target = value.lower()
    return [
        row
        for row in range(2, ws.max_row + 1)
        if norm(ws.cell(row, colmap[field]).value).lower() == target
    ]


def write_non_core_company_audit() -> Path | None:
    if not PRODUCT_MASTER_CSV.exists():
        return None
    AUDIT_DIR.mkdir(parents=True, exist_ok=True)
    with PRODUCT_MASTER_CSV.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))

    by_company: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        if norm(row.get("inclusion_status")).lower() in {"deleted", "excluded"}:
            continue
        by_company[norm(row.get("company"))].append(row)

    audit_rows: list[dict[str, Any]] = []
    for company, company_rows in sorted(by_company.items(), key=lambda item: item[0].lower()):
        if not company_rows:
            continue
        if not all(is_non_core_cosmetic_or_oral_product(row) for row in company_rows):
            continue
        audit_rows.append(
            {
                "company": company,
                "active_product_count_before_scope": len(company_rows),
                "record_ids": "; ".join(norm(row.get("seed_record_id")) for row in company_rows),
                "brands": "; ".join(norm(row.get("brand")) for row in company_rows),
                "product_names": "; ".join(norm(row.get("standard_product_name")) for row in company_rows),
                "commercial_paths": "; ".join(
                    f"{norm(row.get('commercial_path_l1'))} > {norm(row.get('commercial_path_l2'))}"
                    for row in company_rows
                ),
                "scope_action": "exclude_company_from_final_dashboard_results",
                "scope_reason": "only cosmetics/skincare/oral supplements; no qualifying medical-aesthetic device, drug, or equipment product in active rows",
            }
        )

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    audit_path = AUDIT_DIR / f"non_med_aesthetic_company_scope_candidates_{stamp}.csv"
    latest_path = AUDIT_DIR / "non_med_aesthetic_company_scope_candidates_latest.csv"
    fields = [
        "company",
        "active_product_count_before_scope",
        "record_ids",
        "brands",
        "product_names",
        "commercial_paths",
        "scope_action",
        "scope_reason",
    ]
    for path in (audit_path, latest_path):
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            writer.writerows(audit_rows)
    return audit_path


def apply_workbook_cleanup() -> dict[str, Any]:
    if not SOURCE_WORKBOOK.exists():
        raise FileNotFoundError(SOURCE_WORKBOOK)

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = SOURCE_WORKBOOK.with_name(
        f"{SOURCE_WORKBOOK.stem}_backup_before_non_med_aesthetic_scope_cleanup_{stamp}{SOURCE_WORKBOOK.suffix}"
    )
    shutil.copy2(SOURCE_WORKBOOK, backup_path)

    wb = openpyxl.load_workbook(SOURCE_WORKBOOK)
    cleanup_ws = ensure_cleanup_log(wb)
    changes = 0

    product_ws = wb["Product_Lines"]
    product_cols = headers(product_ws)
    target_rows = rows_by_field(product_ws, "Record_ID", TARGET_RECORD_ID)
    if not target_rows:
        raise RuntimeError(f"Could not find {TARGET_RECORD_ID} in Product_Lines")
    for row in target_rows:
        row_id = TARGET_RECORD_ID
        changes += set_cell(product_ws, row, product_cols, "Inclusion_Status", "deleted", cleanup_ws, row_id)
        if "Backfill_Audit" in product_cols:
            changes += set_cell(
                product_ws,
                row,
                product_cols,
                "Backfill_Audit",
                append_note(product_ws.cell(row, product_cols["Backfill_Audit"]).value),
                cleanup_ws,
                row_id,
            )
        changes += set_cell(
            product_ws,
            row,
            product_cols,
            "Verified_Product_Type_CN",
            "口服透明质酸保健品（非医美医疗器械/药品/设备）",
            cleanup_ws,
            row_id,
        )
        changes += set_cell(product_ws, row, product_cols, "Market_Channel", "零售/口服营养补充", cleanup_ws, row_id)
        changes += set_cell(
            product_ws,
            row,
            product_cols,
            "Data_Source",
            append_note(product_ws.cell(row, product_cols["Data_Source"]).value, "manual_scope_cleanup:oral_supplement_not_medical_aesthetics"),
            cleanup_ws,
            row_id,
        )

    if "Companies" in wb.sheetnames:
        company_ws = wb["Companies"]
        company_cols = headers(company_ws)
        for row in rows_by_field(company_ws, "Company", "Biocyte"):
            changes += set_cell(company_ws, row, company_cols, "Product_Count", 0, cleanup_ws, "Biocyte")
            changes += set_cell(company_ws, row, company_cols, "Brand_Count", 0, cleanup_ws, "Biocyte")
            changes += set_cell(company_ws, row, company_cols, "Primary_Track", "", cleanup_ws, "Biocyte")

    for sheet_name, id_field, id_value, fields in [
        ("Product_Master", "seed_record_id", TARGET_RECORD_ID, ["inclusion_status", "Backfill_Audit"]),
        ("Product_Family_Master", "source_record_ids", TARGET_RECORD_ID, ["inclusion_status"]),
        ("Product_SKU_Master", "seed_record_id", TARGET_RECORD_ID, ["inclusion_status", "Backfill_Audit"]),
    ]:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        colmap = headers(ws)
        if id_field not in colmap:
            continue
        for row in range(2, ws.max_row + 1):
            current = norm(ws.cell(row, colmap[id_field]).value)
            if id_value not in {part.strip() for part in current.split(";")} and current != id_value:
                continue
            row_id = norm(ws.cell(row, colmap.get("product_id", colmap.get("product_family_id", colmap.get("sku_id", 1)))).value) or id_value
            for field in fields:
                if field == "Backfill_Audit" and field in colmap:
                    changes += set_cell(ws, row, colmap, field, append_note(ws.cell(row, colmap[field]).value), cleanup_ws, row_id)
                elif field in colmap:
                    changes += set_cell(ws, row, colmap, field, "deleted", cleanup_ws, row_id)
            if "classification_layer" in colmap:
                changes += set_cell(
                    ws,
                    row,
                    colmap,
                    "classification_layer",
                    mark_classification_layer_deleted(ws.cell(row, colmap["classification_layer"]).value),
                    cleanup_ws,
                    row_id,
                )

    wb.save(SOURCE_WORKBOOK)
    return {
        "backup_path": str(backup_path),
        "workbook_path": str(SOURCE_WORKBOOK),
        "changed_cells": changes,
    }


def main() -> None:
    audit_path = write_non_core_company_audit()
    summary = apply_workbook_cleanup()
    summary["audit_path"] = str(audit_path) if audit_path else ""
    summary_path = AUDIT_DIR / "non_med_aesthetic_scope_cleanup_20260615_summary.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
